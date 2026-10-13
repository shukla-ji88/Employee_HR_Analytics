
from pathlib import Path
import sqlite3
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference

# --------------------------------------
# 1. File paths
# --------------------------------------
BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "hr_analytics.db"
REPORT_PATH = BASE_DIR / "HR_Analytics_Final_Report.xlsx"

if not DB_PATH.exists():
    raise FileNotFoundError(
        "Database nahi mili. Pehle python sql_analysis.py run karo."
    )

# --------------------------------------
# 2. Read data from SQLite
# --------------------------------------
with sqlite3.connect(DB_PATH) as conn:

    table_names = set(
        pd.read_sql_query(
            "SELECT name FROM sqlite_master WHERE type='table'",
            conn
        )["name"]
    )

    employees = pd.read_sql_query(
        "SELECT * FROM employee_data ORDER BY EmployeeID",
        conn
    )

    # Enriched employee details using SQL JOIN
    if "department_info" in table_names:
        employee_details = pd.read_sql_query(
            """
            SELECT
                e.EmployeeID,
                e.Name,
                e.Department,
                d.CostCenter,
                d.DepartmentType,
                e.Position,
                e.Salary,
                e.ExperienceYears,
                e.PerformanceRating,
                e.JoiningDate
            FROM employee_data AS e
            LEFT JOIN department_info AS d
                ON e.Department = d.Department
            ORDER BY e.EmployeeID
            """,
            conn
        )
    else:
        employee_details = employees.copy()

    # Lifecycle data may contain simulated exits
    if "employee_lifecycle" in table_names:
        lifecycle = pd.read_sql_query(
            "SELECT * FROM employee_lifecycle ORDER BY EmployeeID",
            conn
        )
    else:
        lifecycle = pd.DataFrame()

if employees.empty:
    raise ValueError("Database mein employee records nahi hain.")

# --------------------------------------
# 3. Department summary
# --------------------------------------
department_summary = (
    employees.groupby("Department", as_index=False)
    .agg(
        Total_Employees=("EmployeeID", "nunique"),
        Average_Salary=("Salary", "mean"),
        Total_Salary_Amount=("Salary", "sum"),
        Average_Experience_Years=("ExperienceYears", "mean"),
        Average_Performance_Rating=("PerformanceRating", "mean")
    )
    .round(2)
)

# --------------------------------------
# 4. Salary band summary
# --------------------------------------
salary = pd.to_numeric(employees["Salary"], errors="coerce")

salary_bands = []

band_definitions = [
    ("Below 35k", salary < 35000),
    ("35k-49,999", (salary >= 35000) & (salary < 50000)),
    ("50k-64,999", (salary >= 50000) & (salary < 65000)),
    ("65k+", salary >= 65000)
]

for band_name, mask in band_definitions:
    salary_bands.append({
        "Salary_Band": band_name,
        "Employee_Count": int(mask.sum()),
        "Total_Salary_Amount": float(salary[mask].sum())
    })

salary_band_summary = pd.DataFrame(salary_bands)

# --------------------------------------
# 5. Hiring trends
# --------------------------------------
joining_dates = pd.to_datetime(
    employees["JoiningDate"],
    errors="coerce"
)

joining_data = employees.loc[joining_dates.notna()].copy()
joining_data["Hire_Date"] = joining_dates.loc[joining_dates.notna()]
joining_data["Hire_Month"] = joining_data["Hire_Date"].dt.strftime("%Y-%m")
joining_data["Hire_Year"] = joining_data["Hire_Date"].dt.strftime("%Y")

monthly_hiring = (
    joining_data.groupby("Hire_Month", as_index=False)
    .agg(New_Hires=("EmployeeID", "nunique"))
    .sort_values("Hire_Month")
)

department_hiring = (
    joining_data.groupby(
        ["Hire_Year", "Department"], as_index=False
    )
    .agg(New_Hires=("EmployeeID", "nunique"))
    .sort_values(["Hire_Year", "Department"])
)

# --------------------------------------
# 6. Lifecycle report
# --------------------------------------
if not lifecycle.empty:
    lifecycle_summary = (
        lifecycle.groupby("EmploymentStatus", as_index=False)
        .agg(Employee_Count=("EmployeeID", "nunique"))
    )

    department_lifecycle = (
        lifecycle.groupby("Department", as_index=False)
        .agg(
            Total_Employees=("EmployeeID", "nunique"),
            Active_Count=(
                "EmploymentStatus",
                lambda values: (values == "Active").sum()
            ),
            Exited_Count=(
                "EmploymentStatus",
                lambda values: (values == "Exited").sum()
            )
        )
    )

    department_lifecycle["Demo_Exit_Share_Percent"] = (
        department_lifecycle["Exited_Count"]
        / department_lifecycle["Total_Employees"].replace(0, pd.NA)
        * 100
    ).fillna(0).round(2)

else:
    lifecycle_summary = pd.DataFrame({
        "Note": [
            "Lifecycle table not found. Run create_lifecycle_demo.py if required."
        ]
    })

    department_lifecycle = pd.DataFrame({
        "Note": ["Lifecycle report unavailable."]
    })

# --------------------------------------
# 7. Executive summary
# --------------------------------------
overview = pd.DataFrame([
    ["Report Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ["Dataset Type", "Sample data for learning and portfolio practice"],
    ["Total Employees", int(employees["EmployeeID"].nunique())],
    ["Total Departments", int(employees["Department"].nunique())],
    ["Average Salary (INR)", round(float(salary.mean()), 2)],
    ["Total Salary Amount (INR)", round(float(salary.sum()), 2)],
    [
        "Average Experience (Years)",
        round(float(employees["ExperienceYears"].mean()), 2)
    ],
    [
        "Average Performance Rating",
        round(float(employees["PerformanceRating"].mean()), 2)
    ],
    [
        "Valid Joining-Date Records",
        int(joining_data["EmployeeID"].nunique())
    ],
    [
        "Exit Data Note",
        "Simulated demo data only; not real company attrition"
        if not lifecycle.empty
        else "Lifecycle table unavailable"
    ]
], columns=["Metric", "Value"])

# --------------------------------------
# 8. Export multiple Excel worksheets
# --------------------------------------
sheets = {
    "Executive Summary": overview,
    "Employee Details": employee_details,
    "Department Summary": department_summary,
    "Salary Bands": salary_band_summary,
    "Hiring Trends": monthly_hiring,
    "Department Hiring": department_hiring,
    "Lifecycle Demo": lifecycle_summary,
    "Department Lifecycle": department_lifecycle
}

with pd.ExcelWriter(REPORT_PATH, engine="openpyxl") as writer:
    for sheet_name, report_df in sheets.items():
        report_df.to_excel(
            writer,
            sheet_name=sheet_name,
            index=False
        )

# --------------------------------------
# 9. Format the workbook professionally
# --------------------------------------
workbook = load_workbook(REPORT_PATH)

header_fill = PatternFill(
    fill_type="solid",
    fgColor="4F46E5"
)

for worksheet in workbook.worksheets:
    worksheet.freeze_panes = "A2"

    if worksheet.max_row > 1:
        worksheet.auto_filter.ref = worksheet.dimensions

    # Format header
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    worksheet.row_dimensions[1].height = 25

    headers = {
        cell.column: str(cell.value or "").lower()
        for cell in worksheet[1]
    }

    # Format numbers
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            if not isinstance(cell.value, (int, float)):
                continue

            header = headers.get(cell.column, "")

            if "salary" in header or "amount" in header:
                cell.number_format = '₹#,##0.00'
            elif "percent" in header or "share" in header:
                cell.number_format = '0.00"%"'
            elif "experience" in header or "rating" in header:
                cell.number_format = "0.00"
            elif "count" in header or "employees" in header or "hires" in header:
                cell.number_format = "#,##0"

    # Adjust column widths
    for column_cells in worksheet.columns:
        column_letter = column_cells[0].column_letter

        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )

        worksheet.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12),
            35
        )

# --------------------------------------
# 10. Add Excel charts
# --------------------------------------
def add_bar_chart(sheet_name, title, data_col, category_col, anchor, y_title):
    worksheet = workbook[sheet_name]

    if worksheet.max_row < 2:
        return

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = title
    chart.y_axis.title = y_title
    chart.x_axis.title = "Category"

    data = Reference(
        worksheet,
        min_col=data_col,
        min_row=1,
        max_row=worksheet.max_row
    )

    categories = Reference(
        worksheet,
        min_col=category_col,
        min_row=2,
        max_row=worksheet.max_row
    )

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 7
    chart.width = 13

    worksheet.add_chart(chart, anchor)


add_bar_chart(
    "Department Summary",
    "Average Salary by Department",
    data_col=3,
    category_col=1,
    anchor="G2",
    y_title="Average Salary (INR)"
)

add_bar_chart(
    "Salary Bands",
    "Employees by Salary Band",
    data_col=2,
    category_col=1,
    anchor="E2",
    y_title="Employee Count"
)

if workbook["Hiring Trends"].max_row > 1:
    worksheet = workbook["Hiring Trends"]

    hiring_chart = LineChart()
    hiring_chart.title = "Monthly Hiring Trend"
    hiring_chart.y_axis.title = "New Hires"
    hiring_chart.x_axis.title = "Month"
    hiring_chart.style = 13

    data = Reference(
        worksheet,
        min_col=2,
        min_row=1,
        max_row=worksheet.max_row
    )

    categories = Reference(
        worksheet,
        min_col=1,
        min_row=2,
        max_row=worksheet.max_row
    )

    hiring_chart.add_data(data, titles_from_data=True)
    hiring_chart.set_categories(categories)
    hiring_chart.height = 7
    hiring_chart.width = 13

    worksheet.add_chart(hiring_chart, "D2")

workbook.save(REPORT_PATH)

print("\nFINAL REPORT CREATED SUCCESSFULLY!")
print("Excel file:", REPORT_PATH.name)
print("Employees included:", employees["EmployeeID"].nunique())
print("Worksheets created:", len(sheets))
print("Report location:", REPORT_PATH)
